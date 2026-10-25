"""Загрузка и проверка xlsx-отчётов бота. Только чтение файлов, сеть не нужна."""

import re
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Iterator, Optional

import pandas as pd
from openpyxl import load_workbook

import settings as S

USER_COLS = ["uid", "name", "username", "n", "replies", "role", "avatar"]
DAILY_COLS = ["uid", "date", "n"]
CHANNEL_COLS = ["uid", "channel_id", "channel", "n"]
HEAT_COLS = ["uid", "weekday", "hour", "n"]
HEAT_TOTAL_COLS = ["weekday", "hour", "n"]
SERVER_COLS = ["uid", "server", "n"]
STAFF_COLS = ["server", "uid", "name", "kind", "roles", "n"]


class LoaderError(Exception):
    """Файл нельзя использовать (не тот формат, повреждён и т. п.)."""


@dataclass
class Dataset:
    path: Path
    title: str
    guild_id: Optional[int]
    is_merged: bool
    schema: int
    full_data: bool                 # есть ли скрытые листы data_* (схема 2)
    tz: str
    period_text: str
    declared_start: Optional[date]
    declared_end: Optional[date]
    data_start: Optional[date]      # фактически первая и последняя даты с сообщениями
    data_end: Optional[date]
    incomplete: bool                # сбор был остановлен вручную
    truncated: bool                 # «По дням (список)» обрезан лимитом
    compare: bool
    meta: dict
    notes: list
    excluded: list
    users: pd.DataFrame
    daily: pd.DataFrame
    prev_daily: pd.DataFrame
    channels: pd.DataFrame
    heat: pd.DataFrame
    heat_total: pd.DataFrame
    servers: pd.DataFrame
    staff: pd.DataFrame
    warnings: list = field(default_factory=list)
    filtered: str = ""

    @property
    def total_messages(self) -> int:
        return int(self.daily["n"].sum()) if len(self.daily) else 0

    @property
    def window(self) -> tuple[Optional[date], Optional[date]]:
        """Календарное окно анализа: заявленный период, иначе фактические даты."""
        return (self.declared_start or self.data_start, self.declared_end or self.data_end)

    @property
    def window_days(self) -> int:
        a, b = self.window
        return (b - a).days + 1 if a and b else 0


# ---------- Мелкие утилиты ----------

def _s(v) -> str:
    return "" if v is None else str(v).strip()


def _to_date(v) -> Optional[date]:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, str):
        t = v.strip()
        for fmt in ("%Y-%m-%d", "%d.%m.%Y"):
            try:
                return datetime.strptime(t, fmt).date()
            except ValueError:
                pass
    return None


def _to_int(v, default: int = 0) -> int:
    if v is None or v == "":
        return default
    try:
        return int(v)
    except (TypeError, ValueError):
        try:
            return int(float(v))
        except (TypeError, ValueError):
            return default


def _uid(v) -> str:
    """ID пользователя как строка из цифр; пустая строка, если это не ID."""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = _s(v)
    return s if s.isdigit() else ""


def _df(rows: list, cols: list) -> pd.DataFrame:
    return pd.DataFrame(rows, columns=cols)


def _rows(wb, name: str) -> Iterator[tuple]:
    if name not in wb.sheetnames:
        return iter(())
    return wb[name].iter_rows(values_only=True)


def _table(wb, name: str) -> list[dict]:
    """Лист «шапка + строки» -> список словарей по именам колонок."""
    it = _rows(wb, name)
    head = next(it, None)
    if not head:
        return []
    head = [_s(h) for h in head]
    return [dict(zip(head, r)) for r in it if r and any(v is not None for v in r)]


# ---------- Разбор листов ----------

def parse_info(wb):
    """«Служебное» -> (одиночные ключи, повторяющиеся ключи, секции-списки)."""
    meta: dict = {}
    multi: dict = defaultdict(list)
    sections: dict = {}
    current = None
    for row in _rows(wb, S.SHEET_INFO):
        row = (list(row) + [None, None])[:2]
        k, v = _s(row[0]), _s(row[1])
        if k and v:
            if k in S.MULTI_KEYS:
                multi[k].append(v)
            else:
                meta.setdefault(k, v)
            current = None
        elif k and not v:                           # заголовок секции: «Нет доступа (3)»
            current = re.sub(r"\s*\(\d+\)\s*$", "", k)
            sections[current] = []
        elif v and current:
            sections[current].append(v)
    return meta, multi, sections


def parse_summary(wb) -> list[dict]:
    """«Сводка». Шапка ищется по тексту: при баннере «сбор остановлен» она сдвигается."""
    out, started = [], False
    for row in _rows(wb, S.SHEET_SUMMARY):
        row = (list(row) + [None] * 10)[:10]
        if not started:
            if _s(row[0]) == "№" and _s(row[1]) == "Пользователь":
                started = True
            continue
        if not isinstance(row[0], (int, float)):
            if any(v is not None for v in row[:4]):
                break                               # итоговая строка «ВСЕГО» или сообщение «не найдено»
            continue
        uid = _uid(row[2])
        if uid:
            out.append({"uid": uid, "display": _s(row[1]), "n": _to_int(row[3]),
                        "role": _s(row[9])})
    return out


def parse_daily(wb) -> pd.DataFrame:
    recs = []
    it = _rows(wb, S.SHEET_DAILY)
    next(it, None)                                  # шапка
    for row in it:
        if not row or len(row) < 4:
            continue
        uid, d, n = _uid(row[1]), _to_date(row[2]), _to_int(row[3])
        if uid and d and n:
            recs.append((uid, d, n))
    df = _df(recs, DAILY_COLS)
    if df.empty:
        df["date"] = pd.to_datetime(df["date"])
        return df
    df["date"] = pd.to_datetime(df["date"])
    return df.groupby(["uid", "date"], as_index=False)["n"].sum()


def parse_staff(wb, default_server: str) -> pd.DataFrame:
    out, server = [], default_server
    for row in _rows(wb, S.SHEET_STAFF):
        row = (list(row) + [None] * 6)[:6]
        a = _s(row[0])
        if a.startswith("Сервер: "):
            server = a[len("Сервер: "):]
            continue
        uid = _uid(row[1])                          # строки-шапки сюда не попадают: «ID» — не число
        kind = S.KIND_BACK.get(_s(row[2]))
        if not uid or kind is None:
            continue
        roles = _s(row[3])
        out.append((server, uid, a, kind, "" if roles == "—" else roles, _to_int(row[4])))
    return _df(out, STAFF_COLS)


def parse_heat_total(wb) -> pd.DataFrame:
    """Запасной вариант для старых файлов: общая тепловая карта с листа «Аналитика»."""
    rows = _rows(wb, S.SHEET_ANALYTICS)
    for row in rows:
        if row and _s(row[0]) == "День \\ час":
            out = []
            for wd in range(7):
                r = next(rows, None)
                if r is None:
                    break
                r = (list(r) + [None] * 25)[:25]
                out += [(wd, hr, _to_int(r[1 + hr])) for hr in range(24)]
            return _df(out, HEAT_TOTAL_COLS)
    return _df([], HEAT_TOTAL_COLS)


# ---------- Сборка набора данных ----------

def load_dataset(path: Path) -> Dataset:
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as e:                           # повреждённый файл, не xlsx и т. п.
        raise LoaderError(f"не удалось открыть файл: {e}")
    try:
        missing = [n for n in S.REQUIRED_SHEETS if n not in wb.sheetnames]
        if missing:
            raise LoaderError("в файле нет листов: " + ", ".join(missing))
        return _build(wb, path)
    finally:
        wb.close()


def _build(wb, path: Path) -> Dataset:
    warnings: list[str] = []
    meta, multi, sections = parse_info(wb)

    title = meta.get("Сервер") or path.stem
    is_merged = (meta.get("Суммарный файл") == "да"
                 or path.name.startswith(S.MERGED_FILE_PREFIX)
                 or title.startswith(S.MERGED_TITLE_PREFIX))
    m = re.search(r"\((\d{15,})\)$", path.stem) or re.match(r"stats_(\d{15,})_", path.stem)
    guild_id = int(m.group(1)) if (m and not is_merged) else None
    schema = _to_int(meta.get("Схема данных"), 1)

    # Период
    d_start = _to_date(meta.get("Период: начало (ISO)"))
    d_end = _to_date(meta.get("Период: конец (ISO)"))
    if d_start is None and d_end is None:            # старые файлы: достаём даты из текста
        found = [_to_date(x) for x in re.findall(r"\d{2}\.\d{2}\.\d{4}", meta.get("Период", ""))]
        if len(found) >= 2 and all(found[:2]):
            d_start, d_end = found[0], found[1]
    gen_date = _to_gen_date(meta.get("Сформирован"))
    last_day = d_end or data_end
    if gen_date and last_day and last_day >= gen_date:
        warnings.append(
            f"Последний день окна ({last_day}) совпадает с днём формирования отчёта и мог быть "
            "неполным: «давность», серии и тренд искажены. Пересоберите отчёт обновлённым ботом."
        )

    attention = multi.get("⚠ Внимание", [])
    incomplete = any("НЕПОЛНЫЕ" in a for a in attention)
    truncated = any("обрезан" in a for a in attention)

    # Основные таблицы
    daily = parse_daily(wb)
    summary = parse_summary(wb)
    have = {n: n in wb.sheetnames for n in (
        S.SHEET_DATA_USERS, S.SHEET_DATA_CHANNELS, S.SHEET_DATA_HEAT,
        S.SHEET_DATA_SERVERS, S.SHEET_DATA_PREV)}
    full = (schema >= 2 and have[S.SHEET_DATA_USERS]
            and have[S.SHEET_DATA_CHANNELS] and have[S.SHEET_DATA_HEAT])

    if full:
        users = _df([
            (_uid(r.get("uid")), _s(r.get("name")), _s(r.get("username")),
             _to_int(r.get("messages")), _to_int(r.get("replies")), _s(r.get("role")),
             _s(r.get("avatar")))
            for r in _table(wb, S.SHEET_DATA_USERS) if _uid(r.get("uid"))
        ], USER_COLS)
        channels = _df([
            (_uid(r.get("uid")), _s(r.get("channel_id")), _s(r.get("channel")), _to_int(r.get("messages")))
            for r in _table(wb, S.SHEET_DATA_CHANNELS) if _uid(r.get("uid"))
        ], CHANNEL_COLS)
        heat = _df([
            (_uid(r.get("uid")), _to_int(r.get("weekday")), _to_int(r.get("hour")), _to_int(r.get("messages")))
            for r in _table(wb, S.SHEET_DATA_HEAT) if _uid(r.get("uid"))
        ], HEAT_COLS)
        heat_total = (heat.groupby(["weekday", "hour"], as_index=False)["n"].sum()
                      if len(heat) else _df([], HEAT_TOTAL_COLS))
    else:
        rows = []
        for r in summary:
            mt = re.match(r"^(.*) \[([^\]]*)\]$", r["display"])
            name, uname = (mt.group(1), mt.group(2)) if mt else (r["display"], "")
            rows.append((r["uid"], name, uname, r["n"], 0, r["role"], ""))
        users = _df(rows, USER_COLS)
        channels = _df([], CHANNEL_COLS)
        heat = _df([], HEAT_COLS)
        heat_total = parse_heat_total(wb)
        warnings.append(
            "Файл создан старой версией бота (нет скрытых листов data_*): нет разбивки по каналам "
            "и личных часов активности, ответы неизвестны. Пересоберите отчёт обновлённым ботом."
        )

    servers = _df([
        (_uid(r.get("uid")), _s(r.get("server")), _to_int(r.get("messages")))
        for r in _table(wb, S.SHEET_DATA_SERVERS) if _uid(r.get("uid"))
    ], SERVER_COLS) if have[S.SHEET_DATA_SERVERS] else _df([], SERVER_COLS)

    prev_rows = []
    if have[S.SHEET_DATA_PREV]:
        for r in _table(wb, S.SHEET_DATA_PREV):
            u, d = _uid(r.get("uid")), _to_date(r.get("date"))
            if u and d:
                prev_rows.append((u, d, _to_int(r.get("messages"))))
    prev_daily = _df(prev_rows, DAILY_COLS)
    prev_daily["date"] = pd.to_datetime(prev_daily["date"])

    staff = parse_staff(wb, title)

    # Фактические границы данных
    data_start = daily["date"].min().date() if len(daily) else None
    data_end = daily["date"].max().date() if len(daily) else None

    # ---------- Проверки ----------
    day_tot = {k: int(v) for k, v in daily.groupby("uid")["n"].sum().items()} if len(daily) else {}

    if daily.empty:
        warnings.append("За выбранный период сообщений нет — анализировать нечего.")
    if incomplete:
        warnings.append("Сбор был остановлен вручную: данные НЕПОЛНЫЕ, все выводы приблизительны.")
    if truncated:
        warnings.append("Лист «По дням (список)» обрезан лимитом строк: суммы по дням занижены.")
    if not meta.get("Часовой пояс"):
        warnings.append("В файле не указан часовой пояс.")
    flt = []
    if meta.get("Фильтр: пользователь", "все") != "все":
        flt.append("пользователь = " + meta["Фильтр: пользователь"])
    if meta.get("Фильтр: канал", "все") != "все":
        flt.append("канал = " + meta["Фильтр: канал"])
    if flt:
        warnings.append("Отчёт построен с фильтром (" + "; ".join(flt) + "): это НЕ весь сервер, "
                        "метрики сервера будут искажены.")

    if not truncated:
        sum_tot = {r["uid"]: r["n"] for r in summary}
        bad = [(u, sum_tot.get(u, 0), day_tot.get(u, 0))
               for u in set(sum_tot) | set(day_tot) if sum_tot.get(u, 0) != day_tot.get(u, 0)]
        if bad:
            u, a, b = bad[0]
            warnings.append(f"«Сводка» и «По дням (список)» расходятся у {len(bad)} польз. "
                            f"(например ID {u}: {a} против {b}).")
        if full:
            ch_tot = ({k: int(v) for k, v in channels.groupby("uid")["n"].sum().items()}
                      if len(channels) else {})
            bad = [u for u in set(ch_tot) | set(day_tot) if ch_tot.get(u, 0) != day_tot.get(u, 0)]
            if bad:
                warnings.append(f"Сумма по каналам не совпадает с суммой по дням у {len(bad)} польз. "
                                f"(например ID {bad[0]}).")

    if data_start and d_start and data_start < d_start:
        warnings.append(f"Есть сообщения раньше заявленного начала периода ({data_start} < {d_start}).")
    if data_end and d_end and data_end > d_end:
        warnings.append(f"Есть сообщения позже заявленного конца периода ({data_end} > {d_end}).")

    return Dataset(
        path=path, title=title, guild_id=guild_id, is_merged=is_merged, schema=schema,
        full_data=full, tz=meta.get("Часовой пояс", ""), period_text=meta.get("Период", ""),
        declared_start=d_start, declared_end=d_end, data_start=data_start, data_end=data_end,
        incomplete=incomplete, truncated=truncated, compare=bool(meta.get("Сравнение с периодом")),
        meta=meta, notes=list(multi.get("Примечание", [])),
        excluded=sections.get("Исключено по правилам (вместе с ветками)", []),
        users=users, daily=daily, prev_daily=prev_daily, channels=channels, heat=heat,
        heat_total=heat_total, servers=servers, staff=staff, warnings=warnings,
        filtered="; ".join(flt),
    )


# ---------- Поиск файлов ----------

def discover(path: Path) -> tuple[list[Path], Optional[Path]]:
    """Возвращает (список xlsx, папка, выбранная автоматически или None).
    Принимает файл, папку с xlsx или родительскую папку output со вложенными stats_*."""
    if path.is_file():
        return ([path] if path.suffix.lower() == ".xlsx" else []), None

    def xlsx_in(folder: Path) -> list[Path]:
        return [p for p in folder.glob("*.xlsx") if not p.name.startswith("~$")]

    files = xlsx_in(path)
    chosen = None
    if not files:
        subs = [d for d in path.glob("stats_*") if d.is_dir() and xlsx_in(d)]
        if subs:
            chosen = max(subs, key=lambda d: d.stat().st_mtime)
            files = xlsx_in(chosen)
    # суммарный файл — последним
    files.sort(key=lambda p: (p.name.startswith(S.MERGED_FILE_PREFIX), p.name.casefold()))
    return files, chosen